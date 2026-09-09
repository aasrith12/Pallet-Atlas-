from pathlib import Path
import re

import openpyxl
import pandas as pd


BASE_DIR = Path(__file__).resolve().parent
OUTPUT_FILE = BASE_DIR / "monte_carlo_combined_clean.csv"
RISK_THRESHOLD = 4.0


WORKBOOKS = [
    {"layer": 1, "path": BASE_DIR / "Monte Carlo Method Layer 1.xlsx", "first_box": 1},
    {"layer": 2, "path": BASE_DIR / "Monte Carlo Method Layer 2.xlsx", "first_box": 6},
    {"layer": 3, "path": BASE_DIR / "Monte Carlo Method Layer 3.xlsx", "first_box": 11},
    {"layer": 4, "path": BASE_DIR / "Monte Carlo Method Layer 4.xlsx", "first_box": 16},
]


def parse_hour(value):
    match = re.fullmatch(r"Hour\s+(\d+)", str(value).strip()) if value is not None else None
    return int(match.group(1)) if match else None


def expected_box_id(workbook_info, sheet_index):
    box_number = workbook_info["first_box"] + sheet_index
    return f"L{workbook_info['layer']}B{box_number}", box_number


def find_stat_row(ws, label):
    for row in ws.iter_rows(min_row=1, max_row=20, min_col=1, max_col=15):
        for cell in row:
            if cell.value == label:
                return cell.row, cell.column
    return None, None


def load_observed_stats(ws):
    average_row, average_col = find_stat_row(ws, "Average")
    _, stdev_col = find_stat_row(ws, "Stdev")
    risk_row, risk_col = find_stat_row(ws, "Risk of Loss (%)")

    if average_row is None or risk_row is None:
        raise ValueError(f"Could not find required summary rows in {ws.title}")

    stats_by_hour = {}
    for col in range(average_col + 1, ws.max_column + 1):
        hour = parse_hour(ws.cell(1, col).value)
        if hour is None:
            continue
        if ws.cell(2, col).value == "Simulation Trial":
            continue

        stats_by_hour[hour] = {
            "observed_average": ws.cell(average_row, col).value,
            "observed_stdev": ws.cell(average_row + 1, col).value if stdev_col is not None else None,
            "monte_carlo_mean": None,
            "monte_carlo_stdev": None,
            "monte_carlo_min": None,
            "monte_carlo_max": None,
            "risk_of_loss_percent": ws.cell(risk_row, col).value if col >= risk_col else None,
        }

    mean_row, mean_col = find_stat_row(ws, "Mean")
    if mean_row is not None:
        for col in range(mean_col + 1, ws.max_column + 1):
            hour = parse_hour(ws.cell(1, col).value)
            if hour is None or hour not in stats_by_hour:
                continue
            if ws.cell(2, col).value == "Simulation Trial":
                continue
            stats_by_hour[hour].update(
                {
                    "monte_carlo_mean": ws.cell(mean_row, col).value,
                    "monte_carlo_stdev": ws.cell(mean_row + 1, col).value,
                    "monte_carlo_min": ws.cell(mean_row + 2, col).value,
                    "monte_carlo_max": ws.cell(mean_row + 3, col).value,
                }
            )

    return stats_by_hour


def simulation_columns(ws):
    columns = []
    for col in range(1, ws.max_column + 1):
        if ws.cell(2, col).value != "Temperature":
            continue
        hour = parse_hour(ws.cell(1, col - 1).value)
        if hour is None:
            continue
        trial_col = col - 1
        columns.append((hour, trial_col, col))
    return columns


def clean_monte_carlo_data():
    records = []

    for workbook_info in WORKBOOKS:
        workbook = openpyxl.load_workbook(workbook_info["path"], data_only=True)
        for sheet_index, ws in enumerate(workbook.worksheets):
            box_id, box_number = expected_box_id(workbook_info, sheet_index)
            source_sheet_name = ws.title
            stats_by_hour = load_observed_stats(ws)

            for hour, trial_col, temp_col in simulation_columns(ws):
                stats = stats_by_hour.get(hour, {})
                for row in range(3, ws.max_row + 1):
                    simulation_trial = ws.cell(row, trial_col).value
                    simulated_value = ws.cell(row, temp_col).value
                    if simulation_trial is None or simulated_value is None:
                        continue
                    if not isinstance(simulated_value, (int, float)):
                        continue

                    records.append(
                        {
                            "source_workbook": workbook_info["path"].name,
                            "source_sheet": source_sheet_name,
                            "box_id": box_id,
                            "box_number": box_number,
                            "layer": workbook_info["layer"],
                            "hour": hour,
                            "simulation_trial": int(simulation_trial),
                            "simulated_value": float(simulated_value),
                            "risk_threshold": RISK_THRESHOLD,
                            "is_risky": float(simulated_value) >= RISK_THRESHOLD,
                            "observed_average": stats.get("observed_average"),
                            "observed_stdev": stats.get("observed_stdev"),
                            "monte_carlo_mean": stats.get("monte_carlo_mean"),
                            "monte_carlo_stdev": stats.get("monte_carlo_stdev"),
                            "monte_carlo_min": stats.get("monte_carlo_min"),
                            "monte_carlo_max": stats.get("monte_carlo_max"),
                            "risk_of_loss_percent": stats.get("risk_of_loss_percent"),
                        }
                    )

    cleaned = pd.DataFrame.from_records(records)
    cleaned = cleaned.sort_values(["layer", "box_number", "hour", "simulation_trial"]).reset_index(drop=True)
    cleaned.to_csv(OUTPUT_FILE, index=False)
    return cleaned


def main():
    cleaned = clean_monte_carlo_data()
    print(f"Wrote {OUTPUT_FILE.name}")
    print(f"Rows: {len(cleaned)}")
    print(f"Boxes: {cleaned['box_id'].nunique()}")
    print(f"Layers: {', '.join(str(layer) for layer in sorted(cleaned['layer'].unique()))}")
    print(f"Hours: {', '.join(str(hour) for hour in sorted(cleaned['hour'].unique()))}")
    print(f"Simulation trials per box-hour: {cleaned.groupby(['box_id', 'hour']).size().min()} to {cleaned.groupby(['box_id', 'hour']).size().max()}")
    print(f"Risk threshold: {RISK_THRESHOLD:g}")


if __name__ == "__main__":
    main()
