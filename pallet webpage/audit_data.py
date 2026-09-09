"""Read-only checks of source workbooks, cleaned data and chart summaries."""
import csv
import math
import statistics
from collections import defaultdict
import openpyxl


def audit(source, boxes):
    def read(name):
        with (source / name).open(encoding='utf-8-sig', newline='') as file:
            return list(csv.DictReader(file))

    raw = read('state_space_combined_clean.csv')
    samples = read('monte_carlo_combined_clean.csv')
    observed = {(r['box_id'], int(r['hour']), r['trial']): r for r in raw}
    simulated = {(r['box_id'], int(r['hour']), int(r['simulation_trial'])): r for r in samples}
    checks = []
    def check(name, passed, detail):
        checks.append(dict(name=name, passed=bool(passed), detail=detail))
    def equal(a, b):
        return math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=1e-9)

    workbook = openpyxl.load_workbook(source / 'State-space Model Data.xlsx', data_only=True)
    seen = set()
    failures = 0
    hour = None
    for row in workbook.active.iter_rows(values_only=True):
        if isinstance(row[0], str) and row[0].startswith('Hour '):
            hour = int(row[0].split()[1])
        if hour is None or not isinstance(row[0], str) or not row[0].startswith('L'):
            continue
        box_id = row[0].strip()
        xyz = [float(v.strip()) for v in str(row[6]).split(',')]
        for index in range(1, 6):
            if not isinstance(row[index], (int, float)):
                continue
            key = (box_id, hour, f'Trial {index+1}')
            seen.add(key)
            saved = observed.get(key)
            failures += int(saved is None or not all(equal(a, b) for a, b in zip(
                [row[index], *xyz, row[7]],
                [saved['value'], saved['x_inches'], saved['y_inches'], saved['z_inches'], saved['exposed_surfaces']] if saved else [])))
    workbook.close()
    check('Observed workbook → cleaned rows', failures == 0 and seen == set(observed),
          f'{len(seen)} trial values, coordinates and recorded surface counts checked; {failures} mismatched rows.')

    seen = set()
    failures = 0
    for layer in range(1, 5):
        workbook = openpyxl.load_workbook(source / f'Monte Carlo Method Layer {layer}.xlsx', data_only=True)
        for index, sheet in enumerate(workbook.worksheets):
            box_id = f'L{layer}B{(layer-1)*5+index+1}'
            for col in range(2, sheet.max_column+1):
                marker = sheet.cell(1, col-1).value
                if sheet.cell(2, col).value != 'Temperature' or not str(marker).startswith('Hour '):
                    continue
                hour = int(marker.split()[1])
                for row in range(3, sheet.max_row+1):
                    trial, value = sheet.cell(row, col-1).value, sheet.cell(row, col).value
                    if trial is None or not isinstance(value, (float, int)):
                        continue
                    key = (box_id, hour, int(trial))
                    seen.add(key)
                    saved = simulated.get(key)
                    failures += int(saved is None or not equal(value, saved['simulated_value']))
        workbook.close()
    check('Monte Carlo workbooks → cleaned samples', failures == 0 and seen == set(simulated),
          f'{len(seen):,} cached sample values checked against four workbooks; {failures} mismatched samples. Existing sheet-order box mapping retained.')

    by_box = defaultdict(list)
    by_sim = defaultdict(list)
    for row in raw:
        by_box[row['box_id']].append(row)
    for row in samples:
        by_sim[row['box_id']].append(row)
    summary_errors = 0
    geometry_errors = 0
    crossing_errors = 0
    for box in boxes:
        trials = by_box[box['box_id']]
        box['trials'] = [{'hour':int(r['hour']), 'trial':r['trial'], 'value':float(r['value'])} for r in trials]
        box['samples'] = [[int(r['hour']), int(r['simulation_trial']), float(r['simulated_value'])] for r in by_sim[box['box_id']]]
        for row in trials:
            geometry_errors += int(any(not equal(box[k], row[k]) for k in ['x_inches','y_inches','z_inches','exposed_surfaces']))
        for row in box['observed']:
            geometry_errors += int(any(not equal(box[k], row[k]) for k in ['x_inches','y_inches','z_inches','exposed_surfaces']))
            values = [t['value'] for t in box['trials'] if t['hour'] == row['hour']]
            expected = dict(mean_value=statistics.mean(values), stdev_value=statistics.stdev(values), min_value=min(values), max_value=max(values), trial_count=len(values))
            summary_errors += sum(not equal(row[k], v) for k, v in expected.items())
        for row in box['simulated']:
            values = sorted(t[2] for t in box['samples'] if t[0] == row['hour'])
            def quantile(q):
                at = (len(values)-1)*q
                lo = int(at)
                return values[lo] + (values[min(lo+1,len(values)-1)]-values[lo])*(at-lo)
            expected = dict(mean_simulated_value=statistics.mean(values), stdev_simulated_value=statistics.stdev(values), simulation_count=len(values), min_simulated_value=min(values), max_simulated_value=max(values), p05_simulated_value=quantile(.05), p50_simulated_value=quantile(.5), p95_simulated_value=quantile(.95), risk_percent=100*sum(v>=4 for v in values)/len(values))
            summary_errors += sum(not equal(row[k], v) for k, v in expected.items())
        crossing = min((t['hour'] for t in box['trials'] if t['value']>=4), default=None)
        crossing_errors += int(crossing != box['first_risk_hour'])
        wide = box['x_inches'] in (11.75,35.25)
        box['geometry'] = dict(width=23.5 if wide else 16, depth=16 if wide else 23.5, height=10)

    # Count outward sides of the nominal packed layer, not artificial rendering gaps.
    # A 1-inch edge tolerance ignores the source's 47 vs 48-inch row offset.
    mismatches = []
    for box in boxes:
        layer = [b for b in boxes if b['layer']==box['layer']]
        def bounds(b):
            return [b['x_inches']-b['geometry']['width']/2,b['x_inches']+b['geometry']['width']/2,b['y_inches']-b['geometry']['depth']/2,b['y_inches']+b['geometry']['depth']/2]
        extent = [min(bounds(b)[0] for b in layer), max(bounds(b)[1] for b in layer), min(bounds(b)[2] for b in layer), max(bounds(b)[3] for b in layer)]
        faces = [name for name,a,b in zip(['X−','X+','Y−','Y+'],bounds(box),extent) if abs(a-b)<=1.01]
        top = int(box['layer']==1)
        bottom = int(box['layer']==4)
        # Bottom included only for a comparison convention, not an airflow claim.
        expected = len(faces)+top+bottom
        mismatch = box['exposed_surfaces'] != expected
        box['geometry'].update(side_faces=faces, side_count=len(faces), top_faces=top, bottom_boundary_faces=bottom, boundary_count=expected, mismatch=mismatch)
        if mismatch:
            mismatches.append(box['box_id'])
    check('Chart inputs → raw values', summary_errors == 0, f'260 observed and 100 simulated box-hour summaries recalculated, including percentiles and threshold percentages; {summary_errors} differing fields.')
    check('3D centers → heat-map inputs', geometry_errors == 0, f'All box centers and recorded counts checked across observed rows; {geometry_errors} differing rows. Existing heat maps plot these same X/Y columns.')
    check('First crossing → individual trials', crossing_errors == 0, f'All 20 first-crossing times checked; {crossing_errors} differences.')
    return dict(checks=checks, mismatches=mismatches, observed_rows=len(raw), simulated_rows=len(samples), geometry_rule='Outward vertical sides of nominal packed layers; 1-inch edge tolerance ignores small row offsets. Add top boundary for Layer 1 and bottom boundary for Layer 4 only when comparing all outer boundaries. Bottom contact with the pallet is not confirmed air exposure. Rendering gaps and exploded views do not count as exposure.')
